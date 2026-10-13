"""Short helpers so a small model can build a real Excel dashboard in about 30 lines instead of 150.

    from forge_office import xl
    wb = xl.workbook()
    data = xl.write_table(wb, "Data", df, formulas={"Revenue": "={Units}*{UnitPrice}"}, formats={"Revenue": "#,##0"})
    by_product = xl.summary(wb, "Summary", "By Product", data, key="Product", keys=products, sums=["Revenue", "Cost"])
    xl.add_formula(by_product, "Profit", "={Revenue}-{Cost}")
    dash = xl.dashboard(wb)
    xl.kpi(dash, 3, "Total revenue", xl.total(data, "Revenue"), "#,##0")
    xl.chart(dash, "bar", "Profit by product", by_product, "Profit", "D2")
    xl.save(wb, "sales.xlsx")

Every number stays a live Excel formula. ``{Name}`` inside a formula is replaced by that column's cell in the same row.
"""

import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
COLORS = {"red": ("FFC7CE", "9C0006"), "green": ("C6EFCE", "006100"), "yellow": ("FFEB9C", "9C5700")}


class Sheet:
    """A written table: where it is and which column letter each header has."""

    def __init__(self, ws, header_row, first, last, cols):
        self.ws, self.header_row, self.first, self.last, self.cols = ws, header_row, first, last, cols

    def col(self, name):
        if name not in self.cols:
            raise KeyError(f"no column {name!r} in {self.ws.title}; columns are {list(self.cols)}")
        return self.cols[name]

    def rng(self, name, absolute=True):
        c = self.col(name)
        return f"'{self.ws.title}'!${c}${self.first}:${c}${self.last}" if absolute else f"{c}{self.first}:{c}{self.last}"


def workbook():
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def _fill(template, cols, row):
    def sub(m):
        name = m.group(1)
        return str(row) if name == "r" else f"{cols[name]}{row}"
    return re.sub(r"\{(\w+)\}", sub, template)


def _style_header(ws, row, ncols, start=1):
    for c in range(start, start + ncols):
        cell = ws.cell(row, c)
        cell.font, cell.fill, cell.alignment = Font(bold=True, color="FFFFFF"), HEAD_FILL, Alignment(horizontal="center")


def autosize(ws, minimum=8, maximum=45):
    for i, col in enumerate(ws.columns, start=1):
        width = max((len(str(c.value)) for c in col if c.value is not None and not str(c.value).startswith("=")), default=minimum)
        ws.column_dimensions[get_column_letter(i)].width = max(minimum, min(maximum, width + 3))


def write_table(wb, name, df, formulas=None, formats=None, as_table=True):
    """Write a DataFrame to a new sheet. ``formulas`` adds computed columns: {"Profit": "={Revenue}-{Cost}"}.
    ``formats`` sets number formats by column name: {"Margin": "0.0%", "Date": "yyyy-mm-dd"}."""
    ws = wb.create_sheet(name)
    headers = list(df.columns) + list(formulas or {})
    cols = {h: get_column_letter(i) for i, h in enumerate(headers, start=1)}
    ws.append(headers)
    for r, row in enumerate(df.itertuples(index=False), start=2):
        values = [v.to_pydatetime() if hasattr(v, "to_pydatetime") else (v.item() if hasattr(v, "item") else v) for v in row]
        values += [_fill(f, cols, r) for f in (formulas or {}).values()]
        ws.append(values)
    last = len(df) + 1
    _style_header(ws, 1, len(headers))
    for h, fmt in (formats or {}).items():
        for c in ws[cols[h]][1:]:
            c.number_format = fmt
    if as_table and len(df):
        t = Table(displayName=re.sub(r"\W", "_", name), ref=f"A1:{get_column_letter(len(headers))}{last}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
    ws.freeze_panes = "A2"
    autosize(ws)
    return Sheet(ws, 1, 2, last, cols)


def summary(wb, sheet, title, data, key, keys, sums, ratio=None, top=None, fmt="#,##0"):
    """A pivot-style block: one row per key with SUMIFS over ``data`` for each column in ``sums``.
    ``ratio=("Margin", "Profit", "Revenue")`` adds a percentage column (needs those columns in ``sums``/added earlier)."""
    ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet)
    top = top or (ws.max_row + 2 if ws.max_row > 1 else 1)
    ws.cell(top, 1, title).font = Font(bold=True, size=12)
    headers = [key] + list(sums)
    for j, h in enumerate(headers, start=1):
        ws.cell(top + 1, j, h)
    _style_header(ws, top + 1, len(headers))
    cols = {h: get_column_letter(j) for j, h in enumerate(headers, start=1)}
    first = top + 2
    for i, k in enumerate(keys, start=first):
        ws.cell(i, 1, k)
        for j, s in enumerate(sums, start=2):
            ws.cell(i, j, f"=SUMIFS({data.rng(s)},{data.rng(key)},$A{i})").number_format = fmt
    block = Sheet(ws, top + 1, first, first + len(keys) - 1, cols)
    if ratio:
        add_formula(block, ratio[0], f"=IF({{{ratio[2]}}}=0,0,{{{ratio[1]}}}/{{{ratio[2]}}})", "0.0%")
    autosize(ws)
    return block


def add_formula(block, name, template, fmt="#,##0"):
    """Add a formula column to a summary block: add_formula(b, "Remaining", "={Budget}-{Actual}")."""
    ws = block.ws
    col = len(block.cols) + 1
    block.cols[name] = get_column_letter(col)
    ws.cell(block.header_row, col, name)
    _style_header(ws, block.header_row, 1, start=col)
    for r in range(block.first, block.last + 1):
        ws.cell(r, col, _fill(template, block.cols, r)).number_format = fmt
    autosize(ws)


def add_values(block, name, values, fmt="#,##0"):
    """Add a column of typed-in inputs (for example the monthly budget of each category)."""
    ws = block.ws
    col = len(block.cols) + 1
    block.cols[name] = get_column_letter(col)
    ws.cell(block.header_row, col, name)
    _style_header(ws, block.header_row, 1, start=col)
    for r, v in zip(range(block.first, block.last + 1), values):
        ws.cell(r, col, v).number_format = fmt
    autosize(ws)


def flag(block, condition, color="red", whole_row=True):
    """Highlight rows of a block where a condition holds: flag(b, "{Remaining}<0")."""
    first = block.first
    cond = re.sub(r"\{(\w+)\}", lambda m: f"${block.cols[m.group(1)]}{first}", condition)
    fill, font = COLORS.get(color, COLORS["red"])
    last_col = get_column_letter(len(block.cols)) if whole_row else block.cols[re.search(r"\{(\w+)\}", condition).group(1)]
    first_col = "A" if whole_row else last_col
    block.ws.conditional_formatting.add(f"{first_col}{first}:{last_col}{block.last}", FormulaRule(
        formula=[cond.lstrip("=")], fill=PatternFill("solid", bgColor=fill), font=Font(color=font)))


def bars(block, column):
    """In-cell data bars on a column."""
    block.ws.conditional_formatting.add(block.rng(column, absolute=False), DataBarRule(start_type="min", end_type="max", color="63BE7B"))


def dashboard(wb, name="Dashboard", title=None):
    ws = wb.create_sheet(name, 0)
    ws["A1"] = title or name
    ws["A1"].font = Font(bold=True, size=16)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 20
    return ws


def total(sheet, column):
    return f"=SUM({sheet.rng(column)})"


def count(sheet, column):
    return f"=COUNTA({sheet.rng(column)})"


def best(block, column, worst=False):
    """Formula giving the label of the row with the highest (or lowest) value in a column of a summary block."""
    fn = "MIN" if worst else "MAX"
    return (f"=INDEX({block.rng(block_key(block), True)},MATCH({fn}({block.rng(column)}),{block.rng(column)},0))")


def block_key(block):
    return next(iter(block.cols))


def kpi(ws, row, label, formula, fmt="General"):
    ws.cell(row, 1, label).font = Font(bold=True)
    ws.cell(row, 2, formula).number_format = fmt


def chart(ws, kind, title, block, column, anchor, width=15, height=7.5):
    """kind: bar, line, pie. Plots one column of a summary block against its labels."""
    ch = {"bar": BarChart, "line": LineChart, "pie": PieChart}[kind]()
    ch.title, ch.width, ch.height = title, width, height
    col = block.col(column)
    idx = list(block.cols).index(column) + 1
    ch.add_data(Reference(block.ws, min_col=idx, min_row=block.header_row, max_row=block.last), titles_from_data=True)
    ch.set_categories(Reference(block.ws, min_col=1, min_row=block.first, max_row=block.last))
    ws.add_chart(ch, anchor)
    return col


def save(wb, path):
    """Save, then open the file again and print what is in it (formulas are computed when Excel opens it)."""
    path = Path(path)
    wb.save(path)
    back = load_workbook(path)
    for ws in back.worksheets:
        n = sum(1 for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
        print(f"saved {path.name} | sheet {ws.title}: {ws.dimensions}, {n} formulas, {len(ws._charts)} charts")
    return path


class _Xl:
    workbook = staticmethod(workbook)
    write_table = staticmethod(write_table)
    summary = staticmethod(summary)
    add_formula = staticmethod(add_formula)
    add_values = staticmethod(add_values)
    flag = staticmethod(flag)
    bars = staticmethod(bars)
    dashboard = staticmethod(dashboard)
    total = staticmethod(total)
    count = staticmethod(count)
    best = staticmethod(best)
    kpi = staticmethod(kpi)
    chart = staticmethod(chart)
    save = staticmethod(save)
    autosize = staticmethod(autosize)


xl = _Xl()

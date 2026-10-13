"""Skills are matched to requests and their recipes run; past successful tasks are saved and recalled."""

import json
import os
import re
import subprocess
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from assistant import experience, skills  # noqa: E402
from assistant.tools import PYLIB  # noqa: E402


def names(text, already=()):
    return [s.name for s in skills.match(text, skills.load_skills(), already)]


def test_requests_pick_the_right_skills(tmp_path, monkeypatch):
    monkeypatch.setenv("FORGE_HOME", str(tmp_path))
    assert names("make an excel sheet of my expenses") == ["excel"]
    assert "pdf" in names("turn this into a PDF invoice")
    assert names("what's the weather in Pune right now?") == ["apis"]
    assert names("convert 100 USD to INR") == ["apis"]
    assert names("make a 5 slide deck about cats") == ["slides"]
    assert names("hello, how was your day?") == []
    assert names("count the words in this paragraph and train the brain") == []    # no 'graph' or 'rain' false hits
    assert names("make an excel sheet", already={"excel"}) == []


def test_user_skill_overrides_builtin(tmp_path, monkeypatch):
    monkeypatch.setenv("FORGE_HOME", str(tmp_path))
    (tmp_path / "skills").mkdir()
    (tmp_path / "skills" / "excel.md").write_text("---\nname: excel\ndescription: mine\ntriggers: excel\n---\nMY NOTES")
    assert skills.load_skills()["excel"].body == "MY NOTES"


def test_every_builtin_recipe_runs(tmp_path):
    for sk in skills.load_skills().values():
        code = "\n".join(re.findall(r"```python\n(.*?)```", sk.body, re.S))
        if not code:
            continue
        env = {**os.environ, "PYTHONPATH": str(PYLIB)}
        run = subprocess.run([sys.executable, "-c", code], cwd=tmp_path, capture_output=True, text=True, timeout=120, env=env)
        assert run.returncode == 0, f"{sk.name}: {run.stderr[-500:]}"
        assert "saved" in run.stdout, sk.name


def tool_turn(task, code, result):
    return [{"role": "user", "content": task + "\n\n(skill notes: excel)\n..."},
            {"role": "assistant", "content": "", "tool_calls": [{"id": "c1", "type": "function",
             "function": {"name": "run_python", "arguments": json.dumps({"code": code})}}]},
            {"role": "tool", "tool_call_id": "c1", "content": result},
            {"role": "assistant", "content": "Done, saved budget.xlsx"}]


def test_working_tasks_are_saved_and_recalled(tmp_path, monkeypatch):
    monkeypatch.setenv("FORGE_HOME", str(tmp_path))
    turn = tool_turn("make an excel budget for my monthly expenses", "wb.save('budget.xlsx')", "saved budget.xlsx")
    assert experience.worked(turn, "Done")
    row = experience.record("make an excel budget for my monthly expenses", turn, "Done")
    assert row["messages"][0]["content"] == "make an excel budget for my monthly expenses"   # notes stripped
    hit = experience.best_match("create an excel budget of expenses for October")
    assert hit and "wb.save" in experience.render(hit)
    assert experience.best_match("tell me a joke about cats") is None
    assert experience.mark_last(good=False) and experience.load() == []


def test_failed_tasks_are_not_saved(tmp_path, monkeypatch):
    monkeypatch.setenv("FORGE_HOME", str(tmp_path))
    bad = tool_turn("make an excel budget", "1/0", "Traceback ...\nZeroDivisionError\n[exit code 1]")
    assert not experience.worked(bad, "Done")
    assert not experience.worked(tool_turn("x", "print(1)", "1"), "(stopped: time limit of 900s reached)")


def test_dashboard_skill_replaces_the_basic_excel_one(tmp_path, monkeypatch):
    monkeypatch.setenv("FORGE_HOME", str(tmp_path))
    assert names("make an excel sales analysis with a pivot table and dashboard") == ["excel_dashboard"]
    assert names("make an excel with my marks") == ["excel"]


def run_recipe(tmp_path, code):
    env = {**os.environ, "PYTHONPATH": str(PYLIB)}
    run = subprocess.run([sys.executable, "-c", code], cwd=tmp_path, capture_output=True, text=True, env=env, timeout=120)
    assert run.returncode == 0, run.stderr[-800:]
    return run.stdout


def test_dashboard_formulas_give_the_same_answers_as_pandas(tmp_path, monkeypatch):
    formulas = pytest.importorskip("formulas")
    monkeypatch.setenv("FORGE_HOME", str(tmp_path))
    code = re.findall(r"```python\n(.*?)```", skills.load_skills()["excel_dashboard"].body, re.S)[0]
    out = run_recipe(tmp_path, code)
    best_product, best_region = re.search(r"best product: (\w+) \| best region: (\w+)", out).groups()
    solved = formulas.ExcelModel().loads(str(tmp_path / "sales_analysis.xlsx")).finish().calculate()
    cell = lambda ref: next(v for k, v in solved.items() if k.upper().endswith(ref)).value[0][0]
    assert cell("DASHBOARD'!B6") == best_product and cell("DASHBOARD'!B7") == best_region
    assert 0.2 < cell("DASHBOARD'!B5") < 0.4


def test_budget_tracker_flags_overspent_categories(tmp_path):
    code = """
import pandas as pd
from forge_office import xl
cats = ["Rent", "Food", "Travel"]
df = pd.DataFrame({"Date": pd.date_range("2026-09-01", periods=6), "Category": ["Rent", "Food", "Food", "Travel", "Food", "Travel"],
                   "Method": ["UPI"] * 6, "Amount": [5000, 1500, 1800, 400, 900, 700]})
wb = xl.workbook()
data = xl.write_table(wb, "Data", df, formats={"Date": "yyyy-mm-dd"})
b = xl.summary(wb, "Budget", "Monthly budget", data, "Category", cats, ["Amount"])
xl.add_values(b, "Budget", [5000, 3000, 1000])
xl.add_formula(b, "Remaining", "={Budget}-{Amount}")
xl.flag(b, "{Remaining}<0")
dash = xl.dashboard(wb)
xl.chart(dash, "bar", "Spend by category", b, "Amount", "D2")
xl.save(wb, "budget.xlsx")
"""
    out = run_recipe(tmp_path, code)
    assert "sheet Budget" in out and "sheet Dashboard" in out and "1 charts" in out
    from openpyxl import load_workbook
    ws = load_workbook(tmp_path / "budget.xlsx")["Budget"]
    assert ws["D3"].value == "=C3-B3"                     # Remaining = Budget - Amount, per category
    assert ws.conditional_formatting                       # overspent rows are highlighted
